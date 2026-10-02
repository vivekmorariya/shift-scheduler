"""
Excel export for Shift Scheduler
Generates a formatted .xlsx file matching the original schedule style
"""

import io
import calendar
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

# Colour palette
COLOURS = {
    "header_bg":    "1F4E79",   # dark blue
    "header_fg":    "FFFFFF",
    "shift1_bg":    "D6E4F0",   # light blue   (Morning)
    "shift2_bg":    "FFF2CC",   # light yellow (Afternoon)
    "shift3_bg":    "FCE4D6",   # light orange (Night)
    "woff_bg":      "E2EFDA",   # light green  (Weekly Off)
    "leave_bg":     "F4CCCC",   # light red    (Leave)
    "eng_row":      "DEEBF7",   # section header – Engineers
    "ops_row":      "FFF2CC",   # section header – Operators
    "tech_row":     "E2EFDA",   # section header – Technicians
    "app_row":      "F4CCCC",   # section header – Apprentices
    "sunday_col":   "D9D9D9",   # light grey column for Sunday
    "alt_sat_col":  "F0F0F0",   # very light grey for alt Saturday
}

ROLE_ORDER = ["Engineers", "Operators", "Technicians", "Apprentices"]
ROLE_COLORS = {
    "Engineers":   COLOURS["eng_row"],
    "Operators":   COLOURS["ops_row"],
    "Technicians": COLOURS["tech_row"],
    "Apprentices": COLOURS["app_row"],
}

SHIFT_FILLS = {
    "1": PatternFill("solid", fgColor=COLOURS["shift1_bg"]),
    "2": PatternFill("solid", fgColor=COLOURS["shift2_bg"]),
    "3": PatternFill("solid", fgColor=COLOURS["shift3_bg"]),
    "W": PatternFill("solid", fgColor=COLOURS["woff_bg"]),
    "L": PatternFill("solid", fgColor=COLOURS["leave_bg"]),
}

THIN = Side(border_style="thin", color="AAAAAA")
THICK = Side(border_style="medium", color="444444")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_BORDER = Border(left=THICK, right=THICK, top=THICK, bottom=THICK)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color="000000", size=9) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size)


def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def get_sundays(year, month):
    cal = calendar.monthcalendar(year, month)
    return {week[6] for week in cal if week[6] != 0}


def get_alt_saturdays(year, month):
    cal = calendar.monthcalendar(year, month)
    sats = [week[5] for week in cal if week[5] != 0]
    return set(sats[i] for i in range(0, len(sats), 2))


def export_to_excel(schedule_result: dict) -> bytes:
    """Return xlsx bytes for the given schedule dict."""
    year = schedule_result["year"]
    month = schedule_result["month"]
    meta = schedule_result["meta"]          # {emp_id: {name, role}}
    emp_scheds = schedule_result["employee_schedules"]   # {emp_id: {day_str: val}}
    days_in_month = schedule_result["days_in_month"]

    days = list(range(1, days_in_month + 1))
    month_label = f"{calendar.month_abbr[month].upper()}. {str(year)[-2:]}"
    sundays = get_sundays(year, month)
    alt_sats = get_alt_saturdays(year, month)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Shift Schedule-{year}"

    # ── Row 1: Title ──────────────────────────────────────────────────────────
    ws.merge_cells("A1:C1")
    ws["A1"] = f"Shift Schedule-{year}"
    ws["A1"].font = _font(bold=True, color=COLOURS["header_fg"], size=13)
    ws["A1"].fill = _fill(COLOURS["header_bg"])
    ws["A1"].alignment = _center()

    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=3 + days_in_month)
    ws.cell(row=1, column=4).value = month_label
    ws.cell(row=1, column=4).font = _font(bold=True, color=COLOURS["header_fg"], size=13)
    ws.cell(row=1, column=4).fill = _fill(COLOURS["header_bg"])
    ws.cell(row=1, column=4).alignment = _center()

    # ── Row 2: Day numbers ────────────────────────────────────────────────────
    ws["A2"] = ""
    ws["B2"] = "Emp ID"
    ws["C2"] = "Emp Name"
    for i, day in enumerate(days, start=4):
        c = ws.cell(row=2, column=i, value=day)
        c.font = _font(bold=True)
        c.alignment = _center()
        c.border = CELL_BORDER
        if day in sundays:
            c.fill = _fill(COLOURS["sunday_col"])
        elif day in alt_sats:
            c.fill = _fill(COLOURS["alt_sat_col"])

    # ── Row 3: Day names ──────────────────────────────────────────────────────
    for i, day in enumerate(days, start=4):
        c = ws.cell(row=3, column=i, value=date(year, month, day).strftime("%a"))
        c.font = _font(bold=False)
        c.alignment = _center()
        c.border = CELL_BORDER
        if day in sundays:
            c.fill = _fill(COLOURS["sunday_col"])
        elif day in alt_sats:
            c.fill = _fill(COLOURS["alt_sat_col"])

    for cell in [ws["A2"], ws["B2"], ws["C2"], ws["A3"], ws["B3"], ws["C3"]]:
        cell.font = _font(bold=True)
        cell.alignment = _center()
        cell.fill = _fill(COLOURS["header_bg"])
        cell.font = _font(bold=True, color=COLOURS["header_fg"])
        cell.border = CELL_BORDER

    # ── Employee rows ─────────────────────────────────────────────────────────
    current_row = 4
    # Build ordered list
    ordered_emps = []
    for role in ROLE_ORDER:
        role_emps = [
            (eid, info) for eid, info in meta.items() if info["role"] == role
        ]
        if role_emps:
            ordered_emps.append(("_section", role))
            ordered_emps.extend(role_emps)

    for item in ordered_emps:
        if item[0] == "_section":
            role = item[1]
            ws.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=3 + days_in_month
            )
            sec_cell = ws.cell(row=current_row, column=1, value=role)
            sec_cell.font = _font(bold=True, size=10)
            sec_cell.fill = _fill(ROLE_COLORS[role])
            sec_cell.alignment = Alignment(horizontal="left", vertical="center")
            sec_cell.border = CELL_BORDER
            current_row += 1
            continue

        eid, info = item
        ws.cell(row=current_row, column=1).border = CELL_BORDER
        ws.cell(row=current_row, column=2, value=eid).alignment = _center()
        ws.cell(row=current_row, column=2).border = CELL_BORDER
        ws.cell(row=current_row, column=3, value=info["name"]).border = CELL_BORDER
        ws.cell(row=current_row, column=3).font = _font(size=9)

        emp_day_map = emp_scheds.get(eid, {})
        for i, day in enumerate(days, start=4):
            val = emp_day_map.get(str(day), "")
            c = ws.cell(row=current_row, column=i, value=val)
            c.alignment = _center()
            c.font = _font(bold=(val in ("W", "L")), size=9)
            c.border = CELL_BORDER
            fill = SHIFT_FILLS.get(val)
            if fill:
                c.fill = fill
            if day in sundays:
                c.fill = _fill(COLOURS["sunday_col"]) if val == "W" else (fill or _fill("FFFFFF"))
            if day in alt_sats and val == "W":
                c.fill = _fill(COLOURS["woff_bg"])

        current_row += 1

    # ── Summary rows ──────────────────────────────────────────────────────────
    current_row += 1
    summary_header = ws.cell(row=current_row, column=1, value="Net persons per day")
    summary_header.font = _font(bold=True)
    summary_header.fill = _fill(COLOURS["header_bg"])
    summary_header.font = Font(bold=True, color=COLOURS["header_fg"])
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    current_row += 1

    for shift_label, shift_key in [("First", "1"), ("Second", "2"), ("Third", "3"), ("W/off", "W")]:
        ws.cell(row=current_row, column=3, value=shift_label).alignment = _center()
        ws.cell(row=current_row, column=3).font = _font(bold=True)
        for i, day in enumerate(days, start=4):
            count = sum(
                1 for eid in emp_scheds
                if emp_scheds[eid].get(str(day)) == shift_key
            )
            c = ws.cell(row=current_row, column=i, value=count)
            c.alignment = _center()
            c.border = CELL_BORDER
        current_row += 1

    # ── Legend ────────────────────────────────────────────────────────────────
    current_row += 1
    legends = [
        ("1 = Morning Shift", COLOURS["shift1_bg"]),
        ("2 = Afternoon Shift", COLOURS["shift2_bg"]),
        ("3 = Night Shift", COLOURS["shift3_bg"]),
        ("W = Weekly Off", COLOURS["woff_bg"]),
        ("L = Leave", COLOURS["leave_bg"]),
    ]
    for j, (label, color) in enumerate(legends):
        col = 1 + j * 2
        ws.merge_cells(start_row=current_row, start_column=col, end_row=current_row, end_column=col + 1)
        c = ws.cell(row=current_row, column=col, value=label)
        c.fill = _fill(color)
        c.font = _font(bold=True, size=8)
        c.alignment = _center()
        c.border = CELL_BORDER

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 28
    for i in range(4, 4 + days_in_month):
        ws.column_dimensions[get_column_letter(i)].width = 4.5

    # Row heights
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 14

    # Freeze panes at D4
    ws.freeze_panes = "D4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
